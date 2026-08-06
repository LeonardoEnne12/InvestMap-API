import io
from typing import Dict
import openpyxl
from app.domain.excel_models import SheetSummary


class ExcelParser:

    @staticmethod
    def _is_valid_data_row(row, header_row=None) -> bool:
        row_dict = {}
        if header_row:
            for idx, col_name in enumerate(header_row):
                val = row[idx] if idx < len(row) else None
                row_dict[col_name] = val

        non_empty_cells = []
        non_empty_text = []

        for cell in row:
            if cell is None:
                continue
            s_val = str(cell).strip()
            if s_val == "":
                continue
            non_empty_cells.append(s_val)

            if not isinstance(cell, (int, float)):
                try:
                    float(s_val.replace(".", "").replace(",", "."))
                except ValueError:
                    non_empty_text.append(s_val)

        if not non_empty_cells:
            return False

        summary_keywords = {
            "total", "total geral", "subtotal", "sub-total",
            "saldo final", "saldo total", "geral", "totais"
        }

        if all(val.lower() in summary_keywords for val in non_empty_cells):
            return False

        if not non_empty_text:
            return False

        if all(val.lower() in summary_keywords for val in non_empty_text):
            return False

        if row_dict:
            primary_cols = {
                "produto", "ativo", "ticker", "descrição", "descricao",
                "asset", "descrição do produto", "descricao do produto", "nome do produto"
            }
            for col_name, val in row_dict.items():
                if col_name.lower() in primary_cols:
                    if val is None:
                        return False
                    s_val = str(val).strip()
                    if s_val == "" or s_val.lower() in summary_keywords:
                        return False
                    break

        return True

    @staticmethod
    def parse_excel_bytes(file_bytes: bytes, preview_limit: int = 5) -> Dict[str, SheetSummary]:
        stream = io.BytesIO(file_bytes)
        wb = openpyxl.load_workbook(stream, data_only=True)
        sheets_summary: Dict[str, SheetSummary] = {}

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))

            if not rows:
                sheets_summary[sheet_name] = SheetSummary(
                    total_rows=0,
                    total_columns=0,
                    columns=[],
                    preview_sample=[]
                )
                continue

            header_row = [str(col) if col is not None else f"Column_{i+1}" for i, col in enumerate(rows[0])]
            data_rows = [r for r in rows[1:] if ExcelParser._is_valid_data_row(r, header_row)]

            preview = []
            for row in data_rows[:preview_limit]:
                row_dict = {}
                for idx, col_name in enumerate(header_row):
                    val = row[idx] if idx < len(row) else None
                    row_dict[col_name] = val
                preview.append(row_dict)

            sheets_summary[sheet_name] = SheetSummary(
                total_rows=len(data_rows),
                total_columns=len(header_row),
                columns=header_row,
                preview_sample=preview
            )

        wb.close()
        return sheets_summary
